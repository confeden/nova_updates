#!/usr/bin/env python3
"""Rebuilds the lists in "список запрещённых" from the official registries.

Lives in confeden/nova_updates as .github/scripts/banned_lists.py and is run by
.github/workflows/publish-banned-lists.yml on a self-hosted runner in Russia:
minjust.gov.ru and fsb.ru do not answer addresses outside the country, so a
GitHub-hosted runner never gets a page.

Each list is written twice, with the same content:

    <name>.json   {"source": ..., "count": ..., "items": [{"n", "name",
                   "short_name", "date"}]}  for programs
    <name>.md     a numbered table                         for people

"date" is the date the entry entered the registry, as the source gives it,
in YYYY-MM-DD. Where the source has no such date for an entry (the oldest
entries of the extremist list) it is the date of the court decision, which
is the date the source shows for that entry.

The Minjust registries of foreign agents and undesirable organisations are
drawn by JavaScript, so every page is opened in headless Chromium. Where the
page offers its "Загрузить реестр" spreadsheet, that file is read; otherwise
the rendered tables or the numbered list on the page are.

A list that comes back empty, smaller than its floor, or much shorter than the
published one is not written: an official page that half-loaded must not wipe
a list someone checks against. The run then fails, so the failure is seen.
"""

import datetime
import io
import json
import os
import re
import sys

from bs4 import BeautifulSoup

OUT_DIR = "список запрещённых"

# (file name, title, page, smallest believable count)
SOURCES = (
    ("иноагенты", "Реестр иностранных агентов",
     "https://minjust.gov.ru/ru/pages/reestr-inostryannykh-agentov/", 300),
    ("террористы", "Единый федеральный список организаций, признанных террористическими",
     "http://www.fsb.ru/fsb/npd/terror.htm", 20),
    ("экстремисты", "Перечень организаций, признанных экстремистскими",
     "https://minjust.gov.ru/ru/documents/7822/", 80),
    ("нежелательные", "Перечень организаций, деятельность которых признана нежелательной",
     "https://minjust.gov.ru/ru/pages/perechen-inostrannyh-i-mezhdunarodnyh-organizacij-deyatelnost-kotoryh-priznana-nezhelatelnoj-na-territorii-rossijskoj-federacii/",
     100),
)

# A new list shorter than this share of the published one is treated as a
# broken page rather than as mass delisting.
MIN_SHARE_OF_PREVIOUS = 0.7

MONTHS = {
    "января": 1, "февраля": 2, "марта": 3, "апреля": 4, "мая": 5, "июня": 6,
    "июля": 7, "августа": 8, "сентября": 9, "октября": 10, "ноября": 11,
    "декабря": 12,
}
DATE_NUM = re.compile(r"\b(\d{1,2})[./](\d{1,2})[./](\d{4}|\d{2})\b")
DATE_ISO = re.compile(r"\b(\d{4})-(\d{2})-(\d{2})\b")
DATE_WORDS = re.compile(r"\b(\d{1,2})\s+(" + "|".join(MONTHS) + r")\s+(\d{4})", re.I)


def log(*args):
    print(*args, flush=True)


# --- dates -------------------------------------------------------------------

def _date(year, month, day):
    year = int(year)
    if year < 100:
        year += 2000
    try:
        return datetime.date(year, int(month), int(day))
    except ValueError:
        return None


def dates_in(text):
    """Every date in the text, in the order it appears."""
    found = []
    for m in DATE_NUM.finditer(text):
        found.append((m.start(), _date(m.group(3), m.group(2), m.group(1))))
    for m in DATE_ISO.finditer(text):
        found.append((m.start(), _date(m.group(1), m.group(2), m.group(3))))
    for m in DATE_WORDS.finditer(text):
        found.append((m.start(), _date(m.group(3), MONTHS[m.group(2).lower()], m.group(1))))
    return [d for _, d in sorted(found, key=lambda x: x[0]) if d]


def cell_date(value):
    if isinstance(value, datetime.datetime):
        return value.date()
    if isinstance(value, datetime.date):
        return value
    found = dates_in(str(value or ""))
    return found[0] if found else None


# --- names -------------------------------------------------------------------

def clean(text):
    text = str(text or "").replace("\xa0", " ").replace("​", "")
    return re.sub(r"\s+", " ", text).strip(" ;,.\t\n")


# Parentheses that describe the decision rather than name the organisation.
DECISION = re.compile(
    r"решени|определени|приговор|суд\b|суда\b|дата размещения|дата включения|"
    r"вступил|вступило|\d{2}\.\d{2}\.\d{4}", re.I)


def split_parens(text):
    """Splits "A (B) C (D)" into the text outside parentheses and the groups."""
    outside, groups, depth, current = [], [], 0, []
    for ch in text:
        if ch == "(":
            if depth == 0:
                current = []
            else:
                current.append(ch)
            depth += 1
        elif ch == ")" and depth:
            depth -= 1
            if depth == 0:
                groups.append("".join(current))
                outside.append("\x00")
            else:
                current.append(ch)
        elif depth:
            current.append(ch)
        else:
            outside.append(ch)
    if depth:  # unbalanced: keep the tail as text
        outside.append("(" + "".join(current))
    return "".join(outside), groups


def name_parts(text):
    """(name, short name) from a list entry such as

        «Исламская группа» («Аль-Гамаа аль-Исламия») (решение ... от 14.02.2003)

    Groups about the decision are dropped. A group that only gives another
    name of the organisation becomes the short name; any other group stays in
    the name, where it belongs to the wording of the source.
    """
    outside, groups = split_parens(clean(text))
    kept, aliases = [], []
    for group in groups:
        group_clean = clean(group)
        if not group_clean or DECISION.search(group_clean):
            kept.append("")
        elif re.match(r"^(далее|сокращ|другие названия|прежн|бывш)", group_clean, re.I) \
                or (group_clean.startswith(("«", '"', "„")) and len(group_clean) < 200):
            aliases.append(re.sub(r"^(далее\s*[-–—]?\s*|сокращ\w*\s*(наименование)?\s*[-–—:]?\s*|другие названия\s*:?\s*)",
                                  "", group_clean, flags=re.I))
            kept.append("")
        else:
            kept.append(f"({group_clean})")
    pieces = outside.split("\x00")
    name = pieces[0]
    for piece, group in zip(pieces[1:], kept):
        name += (" " + group if group else "") + piece
    return clean(name), "; ".join(a for a in aliases if a)


# --- tables ------------------------------------------------------------------

HEADER_NAME = re.compile(r"наименовани|ф\.?\s*и\.?\s*о|фамилия|организаци", re.I)


def find_columns(header):
    """Column indices for name, short name, entry date and exclusion date."""
    cols = {"name": None, "short": None, "date": None, "excluded": None}
    date_rank = None
    for i, cell in enumerate(header):
        h = clean(cell).lower()
        if not h:
            continue
        if "исключ" in h:
            cols["excluded"] = i
            continue
        if "сокращ" in h and cols["short"] is None:
            cols["short"] = i
        elif HEADER_NAME.search(h) and "иностран" not in h and cols["name"] is None:
            cols["name"] = i
        if "дата" in h:
            rank = 0 if re.search(r"включ|внесен", h) else 1 if "решени" in h else 2
            if date_rank is None or rank < date_rank:
                cols["date"], date_rank = i, rank
    return cols


def table_items(rows):
    """Items from a table given as a list of rows of cell values."""
    rows = [list(r) for r in rows if any(clean(c) for c in r)]
    if not rows:
        return []
    header_at = next((i for i, r in enumerate(rows[:20])
                      if any(HEADER_NAME.search(clean(c)) for c in r if isinstance(c, str))), None)
    cols = find_columns(rows[header_at]) if header_at is not None else {}
    log(f"    header row {header_at}: {[clean(c) for c in rows[header_at]] if header_at is not None else None}")
    log(f"    columns: {cols}")
    body = rows[header_at + 1:] if header_at is not None else rows

    items = []
    for row in body:
        cells = [clean(c) if not isinstance(c, (datetime.date, datetime.datetime)) else c for c in row]
        if cols.get("excluded") is not None and cols["excluded"] < len(row) and cell_date(row[cols["excluded"]]):
            continue
        if cols.get("name") is not None and cols["name"] < len(cells):
            raw_name = cells[cols["name"]]
        else:
            texts = [c for c in cells if isinstance(c, str) and not re.fullmatch(r"[\d.\s№]*", c)]
            raw_name = max(texts, key=len) if texts else ""
        if not raw_name or HEADER_NAME.fullmatch(raw_name):
            continue
        name, short = name_parts(raw_name)
        if cols.get("short") is not None and cols["short"] < len(cells) and cells[cols["short"]]:
            short = cells[cols["short"]]
        date = None
        if cols.get("date") is not None and cols["date"] < len(row):
            date = cell_date(row[cols["date"]])
        if date is None:
            # No date column: the first date anywhere else in the row, which
            # on these pages is the decision that put the entry on the list.
            for i, c in enumerate(row):
                if i != cols.get("name"):
                    date = cell_date(c)
                    if date:
                        break
            if date is None:
                found = dates_in(raw_name)
                date = found[0] if found else None
        if name:
            items.append({"name": name, "short_name": short, "date": date})
    return items


def html_tables(soup):
    tables = []
    for table in soup.find_all("table"):
        rows = []
        for tr in table.find_all("tr"):
            rows.append([td.get_text(" ", strip=True) for td in tr.find_all(["td", "th"])])
        if len(rows) > 1:
            tables.append(rows)
    return tables


def sheet_rows(data, filename):
    """Rows of the first sheet of an .xlsx or .xls file."""
    if data[:2] == b"PK":
        import openpyxl
        book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sheet = book.worksheets[0]
        return [list(r) for r in sheet.iter_rows(values_only=True)]
    if data[:4] == b"\xd0\xcf\x11\xe0":
        import xlrd
        book = xlrd.open_workbook(file_contents=data)
        sheet = book.sheet_by_index(0)
        rows = []
        for r in range(sheet.nrows):
            row = []
            for c in range(sheet.ncols):
                cell = sheet.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    row.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode))
                else:
                    row.append(cell.value)
            rows.append(row)
        return rows
    # Some "XLS" downloads are an HTML table with an .xls name.
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError:
        text = data.decode("cp1251", "replace")
    if "<table" in text.lower():
        tables = html_tables(BeautifulSoup(text, "html.parser"))
        return max(tables, key=len) if tables else []
    raise ValueError(f"unknown spreadsheet format: {filename}")


# --- numbered text lists -----------------------------------------------------

ENTRY_START = re.compile(r"^\s*(\d{1,4})\s*[.)]\s+(\S.*)$")


def numbered_items(soup):
    """Items from a list written as "1. ...", "2. ..." paragraphs.

    Takes the longest run of lines numbered 1, 2, 3, ... and joins each
    entry with the lines after it up to the next number. Lines after the last
    entry are only taken while they look like its decision note, so the page
    footer does not end up in the last name.
    """
    for tag in soup(["script", "style", "noscript", "header", "footer", "nav"]):
        tag.decompose()
    lines = [clean(l) for l in soup.get_text("\n").split("\n")]
    lines = [l for l in lines if l]

    best = []
    for start, line in enumerate(lines):
        m = ENTRY_START.match(line)
        if not m or m.group(1) != "1":
            continue
        run, expected, i = [], 1, start
        while i < len(lines):
            m = ENTRY_START.match(lines[i])
            if m and int(m.group(1)) == expected:
                run.append([m.group(2)])
                expected += 1
            elif m and int(m.group(1)) > expected and run:
                break
            elif run:
                run[-1].append(lines[i])
            i += 1
        if len(run) > len(best):
            best = run
    if not best:
        return []
    last = best[-1]
    best[-1] = [last[0]] + [l for l in last[1:4] if l.startswith("(") or l[:1].islower()]

    items = []
    for parts in best:
        text = " ".join(parts)
        posted = re.search(r"дата\s+(размещения|включения)[^:]*:\s*([\d.]+)", text, re.I)
        date = cell_date(posted.group(2)) if posted else None
        if date is None:
            found = dates_in(text)
            date = found[0] if found else None
        name, short = name_parts(text)
        if name:
            items.append({"name": name, "short_name": short, "date": date})
    return items


# --- fetching ----------------------------------------------------------------

def fetch(browser, url):
    """The rendered page and, if the page offers one, its spreadsheet."""
    page = browser.new_page(locale="ru-RU", accept_downloads=True)
    try:
        page.goto(url, wait_until="domcontentloaded", timeout=120_000)
        try:
            page.wait_for_load_state("networkidle", timeout=60_000)
        except Exception:
            log("    page kept loading; reading what is there")
        log(f"    title: {page.title()!r}")
        sheet = None
        button = page.locator("a:has-text('Загрузить реестр'), button:has-text('Загрузить реестр'),"
                              " [title*='XLS'], a[href$='.xlsx'], a[href$='.xls']").first
        if button.count():
            try:
                with page.expect_download(timeout=90_000) as info:
                    button.click()
                download = info.value
                with open(download.path(), "rb") as fh:
                    sheet = (fh.read(), download.suggested_filename)
                log(f"    downloaded {sheet[1]!r}, {len(sheet[0])} bytes")
            except Exception as error:
                log(f"    no spreadsheet from the page: {error.__class__.__name__}: {error}")
        return page.content(), sheet
    finally:
        page.close()


def parse(html, sheet):
    if sheet:
        try:
            rows = sheet_rows(*sheet)
            log(f"    spreadsheet rows: {len(rows)}; first: {[clean(c) for c in rows[0]] if rows else None}")
            items = table_items(rows)
            if items:
                return items, "spreadsheet"
        except Exception as error:
            log(f"    spreadsheet unreadable: {error}")
    soup = BeautifulSoup(html, "html.parser")
    tables = html_tables(soup)
    log(f"    tables on page: {[len(t) for t in tables]}")
    from_table = table_items(max(tables, key=len)) if tables else []
    from_list = numbered_items(soup)
    if len(from_table) >= len(from_list):
        return from_table, "table"
    return from_list, "numbered list"


# --- writing -----------------------------------------------------------------

def published_count(path):
    try:
        with open(path, encoding="utf-8") as fh:
            return json.load(fh).get("count", 0)
    except (OSError, ValueError):
        return 0


def dedupe(items):
    seen, out = set(), []
    for item in items:
        key = (item["name"].lower(), item["date"])
        if key not in seen:
            seen.add(key)
            out.append(item)
    return out


def md_cell(text):
    return str(text or "").replace("|", "\\|")


def write(key, title, url, items):
    os.makedirs(OUT_DIR, exist_ok=True)
    numbered = [{"n": i, "name": it["name"], "short_name": it["short_name"],
                 "date": it["date"].isoformat() if it["date"] else None}
                for i, it in enumerate(items, 1)]
    with open(os.path.join(OUT_DIR, key + ".json"), "w", encoding="utf-8") as fh:
        json.dump({"source": url, "count": len(numbered), "items": numbered},
                  fh, ensure_ascii=False, indent=1)
        fh.write("\n")
    lines = [f"# {title}", "", f"Источник: {url}", "", f"Записей: {len(numbered)}", "",
             "| № | Название | Сокращённое | Дата внесения |", "|---|---|---|---|"]
    for it in numbered:
        date = (datetime.date.fromisoformat(it["date"]).strftime("%d.%m.%Y") if it["date"] else "")
        lines.append(f"| {it['n']} | {md_cell(it['name'])} | {md_cell(it['short_name'])} | {date} |")
    with open(os.path.join(OUT_DIR, key + ".md"), "w", encoding="utf-8") as fh:
        fh.write("\n".join(lines) + "\n")


def main():
    from playwright.sync_api import sync_playwright

    only = set(sys.argv[1:])
    failed = []
    with sync_playwright() as pw:
        chromium = os.environ.get("CHROMIUM_PATH") or None
        browser = pw.chromium.launch(executable_path=chromium, headless=True)
        for key, title, url, floor in SOURCES:
            if only and key not in only:
                continue
            log(f"== {key}: {url}")
            try:
                html, sheet = fetch(browser, url)
                items, how = parse(html, sheet)
            except Exception as error:
                log(f"    failed: {error.__class__.__name__}: {error}")
                failed.append(key)
                continue
            items = dedupe(items)
            log(f"    {len(items)} entries from the {how}")
            for it in items[:3] + items[-2:]:
                log(f"      {it}")
            previous = published_count(os.path.join(OUT_DIR, key + ".json"))
            if len(items) < floor or len(items) < previous * MIN_SHARE_OF_PREVIOUS:
                log(f"    refusing to publish: {len(items)} entries, floor {floor}, published {previous}")
                failed.append(key)
                continue
            write(key, title, url, items)
        browser.close()
    if failed:
        log(f"not updated: {', '.join(failed)}")
        sys.exit(1)


if __name__ == "__main__":
    main()
